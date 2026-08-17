import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy import or_
from sqlalchemy.orm import Session
from ..database import get_db
from ..models import DireccionInmueble, User, UploadLog
from ..services.auth import get_current_user, require_superadmin

router = APIRouter(prefix="/api/direcciones", tags=["direcciones"])


import re

def _norm(s):
    return (s or "").strip().lower()


def _clean_jurisdiccion(s):
    """'11. COYHAIQUE' -> 'Coyhaique' (quita prefijo numérico, capitaliza)."""
    s = (s or "").strip()
    s = re.sub(r"^\d+\.\s*", "", s)
    return s.title() if s else s


@router.post("/upload")
def upload_direcciones(
    file: UploadFile = File(...),
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx")

    content = file.file.read()
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el Excel: {e}")

    header_row_idx = None
    col_map = {}
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        cells = [_norm(str(c)) if c is not None else "" for c in row]
        if any("direccion seal" in c for c in cells):
            header_row_idx = i
            for idx, c in enumerate(cells):
                if "jurisdic" in c:
                    col_map["jurisdiccion"] = idx
                elif "nombre edificio" in c:
                    col_map["nombre_edificio"] = idx
                elif "direccion seal" in c:
                    col_map["direccion_seal"] = idx
                elif "situaci" in c:
                    col_map["situacion_actual"] = idx
            break

    if header_row_idx is None or "direccion_seal" not in col_map or "nombre_edificio" not in col_map:
        raise HTTPException(
            status_code=400,
            detail="No se encontraron las columnas esperadas (Nombre Edificio, Direccion SEAL) en el Excel",
        )

    nuevas = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        def get(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return ""
            val = row[idx]
            return str(val).strip() if val is not None else ""

        nombre_edificio = get("nombre_edificio")
        direccion_seal = get("direccion_seal")
        if not nombre_edificio or not direccion_seal:
            continue
        nuevas.append(
            DireccionInmueble(
                jurisdiccion=_clean_jurisdiccion(get("jurisdiccion")) or "Sin jurisdicción",
                nombre_edificio=nombre_edificio,
                direccion_seal=direccion_seal,
                situacion_actual=get("situacion_actual") or "",
            )
        )

    if not nuevas:
        raise HTTPException(status_code=400, detail="No se encontraron filas válidas para importar")

    db.query(DireccionInmueble).delete()
    db.add_all(nuevas)
    db.add(UploadLog(
        source="direcciones",
        filename=file.filename,
        record_count=len(nuevas),
        uploaded_by=getattr(admin, "email", "") or getattr(admin, "name", ""),
    ))
    db.commit()

    return {"message": f"{len(nuevas)} direcciones importadas correctamente"}


@router.get("/jurisdicciones")
def list_jurisdicciones(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = (
        db.query(DireccionInmueble.jurisdiccion)
        .filter(DireccionInmueble.situacion_actual.ilike("en uso"))
        .distinct()
        .order_by(DireccionInmueble.jurisdiccion)
        .all()
    )
    return [r[0] for r in rows]


@router.get("/buscar")
def buscar_direcciones(
    jurisdiccion: str,
    q: str = "",
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(DireccionInmueble).filter(
        DireccionInmueble.jurisdiccion == jurisdiccion,
        DireccionInmueble.situacion_actual.ilike("en uso"),
    )
    if q:
        like = f"%{q}%"
        query = query.filter(
            or_(
                DireccionInmueble.nombre_edificio.ilike(like),
                DireccionInmueble.direccion_seal.ilike(like),
            )
        )
    results = query.order_by(DireccionInmueble.nombre_edificio).limit(20).all()
    return [
        {"id": r.id, "nombre_edificio": r.nombre_edificio, "direccion_seal": r.direccion_seal}
        for r in results
    ]


@router.get("/count")
def count_direcciones(admin: User = Depends(require_superadmin), db: Session = Depends(get_db)):
    return {"total": db.query(DireccionInmueble).count()}
