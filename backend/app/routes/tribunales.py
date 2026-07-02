import io
import time
import requests
from fastapi import APIRouter, Depends, HTTPException
from ..models import User
from ..services.auth import get_current_user

router = APIRouter(prefix="/api/tribunales", tags=["tribunales"])

ONEDRIVE_URL = (
    "https://pjudcl-my.sharepoint.com/:x:/g/personal/ccruzatv_poderjudicial_cl"
    "/IQBMv76Z0109S5gOsdCEQf8wAeqRBtru4ECh7wB2dcsKTcw?e=HDE9Cg&download=1"
)

_cache: dict = {"wb": None, "ts": 0.0}
_CACHE_TTL = 3600  # 1 hora


def _load_workbook():
    import openpyxl
    now = time.time()
    if _cache["wb"] is not None and now - _cache["ts"] < _CACHE_TTL:
        return _cache["wb"]
    try:
        r = requests.get(ONEDRIVE_URL, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
        r.raise_for_status()
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=False, data_only=True)
        _cache["wb"] = wb
        _cache["ts"] = now
        return wb
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"No se pudo cargar el Excel de OneDrive: {e}")


def get_workbook():
    return _load_workbook()


@router.get("/")
def list_jurisdicciones(user: User = Depends(get_current_user)):
    """Returns the sheet names (jurisdictions) available in the OneDrive Excel."""
    wb = _load_workbook()
    return wb.sheetnames


@router.post("/refresh")
def refresh_cache(user: User = Depends(get_current_user)):
    """Force reload of the Excel from OneDrive."""
    if user.role not in ("admin", "superadmin"):
        raise HTTPException(status_code=403, detail="Sin permisos")
    _cache["wb"] = None
    _cache["ts"] = 0.0
    _load_workbook()
    return {"message": "Caché actualizado"}
