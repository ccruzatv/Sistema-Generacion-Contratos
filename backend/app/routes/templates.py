import io
import os
import uuid
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Template, TemplateField, User
from ..schemas import (
    TemplateResponse,
    TemplateListResponse,
    TemplateUpdate,
    TemplateFieldUpdate,
    TemplateFieldCreate,
)
from ..services.template_parser import extract_placeholders
from ..services.auth import get_current_user, require_superadmin

router = APIRouter(prefix="/api/templates", tags=["templates"])

TEMPLATES_DIR = os.environ.get(
    "TEMPLATES_DIR",
    os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "templates"),
)
os.makedirs(TEMPLATES_DIR, exist_ok=True)


def _resolve_path(stored: str) -> str:
    """Return absolute path for a stored file_path value (supports legacy absolute paths)."""
    if os.path.isabs(stored):
        return stored
    return os.path.join(TEMPLATES_DIR, stored)


def _to_list_response(t: Template) -> TemplateListResponse:
    return TemplateListResponse(
        id=t.id,
        name=t.name,
        description=t.description,
        category_id=t.category_id,
        category_name=t.category.name if t.category else None,
        created_at=t.created_at,
        field_count=len(t.fields),
    )


def _to_detail_response(t: Template) -> TemplateResponse:
    resp = TemplateResponse.model_validate(t)
    resp.category_name = t.category.name if t.category else None
    return resp


@router.get("/", response_model=list[TemplateListResponse])
def list_templates(
    category_id: int | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(Template)
    if category_id is not None:
        query = query.filter(Template.category_id == category_id)
    templates = query.order_by(Template.created_at.desc()).all()
    return [_to_list_response(t) for t in templates]


@router.get("/{template_id}", response_model=TemplateResponse)
def get_template(
    template_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")
    return _to_detail_response(template)


@router.post("/", response_model=TemplateResponse)
def create_template(
    name: str = Form(...),
    description: str = Form(""),
    category_id: int = Form(None),
    file: UploadFile = File(...),
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    if not file.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .docx")

    content = file.file.read()
    MAX_SIZE = 100 * 1024 * 1024  # 100 MB
    if len(content) > MAX_SIZE:
        raise HTTPException(status_code=400, detail="El archivo supera el tamaño máximo permitido (100 MB)")
    # Validate DOCX magic bytes (ZIP: PK\x03\x04)
    if len(content) < 4 or content[:4] != b"PK\x03\x04":
        raise HTTPException(status_code=400, detail="El archivo no es un .docx válido")
    # Validate DOCX structure
    try:
        from docx import Document as _DocxDocument
        _DocxDocument(io.BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo .docx está corrupto o no es válido")

    file_id = uuid.uuid4().hex[:12]
    filename = f"{file_id}_{file.filename}"
    file_path = os.path.join(TEMPLATES_DIR, filename)

    with open(file_path, "wb") as f:
        f.write(content)

    try:
        detected_fields = extract_placeholders(file_path)
    except Exception as e:
        os.remove(file_path)
        raise HTTPException(
            status_code=400, detail=f"Error al procesar la plantilla: {str(e)}"
        )

    template = Template(
        name=name,
        description=description,
        file_path=filename,   # store only filename for Docker portability
        category_id=category_id,
    )
    db.add(template)
    db.flush()

    for field_data in detected_fields:
        field = TemplateField(template_id=template.id, **field_data)
        db.add(field)

    db.commit()
    db.refresh(template)
    return _to_detail_response(template)


@router.put("/{template_id}", response_model=TemplateResponse)
def update_template(
    template_id: int,
    update: TemplateUpdate,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")

    if update.name is not None:
        template.name = update.name
    if update.description is not None:
        template.description = update.description
    if update.category_id is not None:
        template.category_id = update.category_id

    db.commit()
    db.refresh(template)
    return _to_detail_response(template)


@router.delete("/{template_id}")
def delete_template(
    template_id: int,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")

    resolved = _resolve_path(template.file_path)
    if os.path.exists(resolved):
        os.remove(resolved)

    db.delete(template)
    db.commit()
    return {"message": "Plantilla eliminada exitosamente"}


@router.post("/{template_id}/reparse", response_model=TemplateResponse)
def reparse_template(
    template_id: int,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")

    detected_fields = extract_placeholders(_resolve_path(template.file_path))
    db.query(TemplateField).filter(TemplateField.template_id == template_id).delete()

    for field_data in detected_fields:
        field = TemplateField(template_id=template.id, **field_data)
        db.add(field)

    db.commit()
    db.refresh(template)
    return _to_detail_response(template)


@router.post("/reparse-all")
def reparse_all_templates(
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    templates = db.query(Template).all()
    ok, failed = [], []

    for template in templates:
        try:
            detected_fields = extract_placeholders(_resolve_path(template.file_path))
            db.query(TemplateField).filter(TemplateField.template_id == template.id).delete()
            for field_data in detected_fields:
                db.add(TemplateField(template_id=template.id, **field_data))
            ok.append(template.name)
        except Exception as e:
            failed.append({"name": template.name, "error": str(e)})

    db.commit()
    return {
        "message": f"{len(ok)} plantillas re-detectadas correctamente" + (f", {len(failed)} con errores" if failed else ""),
        "updated": ok,
        "failed": failed,
    }


@router.put("/{template_id}/fields/{field_id}", response_model=TemplateResponse)
def update_field(
    template_id: int,
    field_id: int,
    update: TemplateFieldUpdate,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    field = (
        db.query(TemplateField)
        .filter(TemplateField.id == field_id, TemplateField.template_id == template_id)
        .first()
    )
    if not field:
        raise HTTPException(status_code=404, detail="Campo no encontrado")

    for key, value in update.model_dump(exclude_unset=True).items():
        setattr(field, key, value)

    db.commit()
    template = db.query(Template).filter(Template.id == template_id).first()
    db.refresh(template)
    return _to_detail_response(template)


@router.post("/{template_id}/fields", response_model=TemplateResponse)
def add_field(
    template_id: int,
    field_data: TemplateFieldCreate,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    template = db.query(Template).filter(Template.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")

    field = TemplateField(template_id=template_id, **field_data.model_dump())
    db.add(field)
    db.commit()
    db.refresh(template)
    return _to_detail_response(template)


@router.delete("/{template_id}/fields/{field_id}")
def delete_field(
    template_id: int,
    field_id: int,
    admin: User = Depends(require_superadmin),
    db: Session = Depends(get_db),
):
    field = (
        db.query(TemplateField)
        .filter(TemplateField.id == field_id, TemplateField.template_id == template_id)
        .first()
    )
    if not field:
        raise HTTPException(status_code=404, detail="Campo no encontrado")

    db.delete(field)
    db.commit()
    return {"message": "Campo eliminado exitosamente"}
