"""
Ejecutar desde la carpeta backend:
    python inspect_excel.py

Requiere: pip install requests openpyxl
"""
import io
import requests
import openpyxl
from openpyxl.utils import get_column_letter

ONEDRIVE_URL = (
    "https://pjudcl-my.sharepoint.com/:x:/g/personal/ccruzatv_poderjudicial_cl"
    "/IQBMv76Z0109S5gOsdCEQf8wAeqRBtru4ECh7wB2dcsKTcw?e=HDE9Cg&download=1"
)

print("Descargando Excel desde OneDrive...")
r = requests.get(ONEDRIVE_URL, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
r.raise_for_status()
wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=False, data_only=True)
print(f"Hojas encontradas ({len(wb.sheetnames)}): {wb.sheetnames}\n")

PAGE_WIDTH_TWIPS = 9072

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print("=" * 60)
    print(f"HOJA: {sheet_name}")
    print(f"  Dimensiones: {ws.max_row} filas x {ws.max_column} cols")

    # Columnas y anchos
    print(f"  Anchos de columna (Excel chars → twips normalizados):")
    raw_widths = []
    for ci in range(ws.max_column or 0):
        col_letter = get_column_letter(ci + 1)
        dim = ws.column_dimensions.get(col_letter)
        w = dim.width if dim and dim.width else 8.43
        raw_widths.append(w)
    total = sum(raw_widths) or 1
    MIN_COL = max(300, int(PAGE_WIDTH_TWIPS * 0.05))
    clamped = [max(MIN_COL, int(w / total * PAGE_WIDTH_TWIPS)) for w in raw_widths]
    ct = sum(clamped)
    final = [max(MIN_COL, int(t / ct * PAGE_WIDTH_TWIPS)) for t in clamped]
    for ci, (rw, fw) in enumerate(zip(raw_widths, final)):
        col_letter = get_column_letter(ci + 1)
        print(f"    Col {col_letter}: {rw:.2f} chars → {fw} twips ({fw/567:.2f} cm)")

    # Celdas fusionadas
    merges = list(ws.merged_cells.ranges)
    print(f"  Celdas fusionadas ({len(merges)}):")
    for mr in merges:
        print(f"    {mr} → filas {mr.min_row}-{mr.max_row}, cols {mr.min_col}-{mr.max_col}")

    # Primeras filas de datos
    print(f"  Primeras 3 filas de datos:")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 3:
            break
        vals = [str(c)[:30] if c is not None else "(vacío)" for c in row]
        print(f"    Fila {i+1}: {vals}")
    print()

print("Inspección completa.")
